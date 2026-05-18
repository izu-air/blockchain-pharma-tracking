#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
embed_diagrams.py
Replaces [МЕСТО ДЛЯ РИСУНКА X.X] stubs in Диплом.docx with actual PNG images.
Run: python embed_diagrams.py
"""
from docx import Document
from docx.shared import Inches, Cm, Pt
from docx.oxml.ns import qn
import os
import re

# Strict number extraction so "2.1" stub does not eat "2.10" / "2.12".
STUB_RE = re.compile(r'РИСУНКА\s+(\d+\.\d+)', re.IGNORECASE)

INPUT  = 'Диплом.docx'
OUTPUT = 'Диплом.docx'

FIG_MAP = {
    '2.1':  ('diagrams/fig_2_1_usecase.png',       Inches(6.2)),
    '2.2':  ('diagrams/fig_2_2_conceptual_er.png',  Inches(6.0)),
    '2.3':  ('diagrams/fig_2_3_logical_er.png',     Inches(6.2)),
    '2.4':  ('diagrams/fig_2_4_physical_er.png',    Inches(6.2)),
    '2.5':  ('diagrams/fig_2_5_component.png',      Inches(6.0)),
    '2.6':  ('diagrams/fig_2_6_seq_register.png',   Inches(6.2)),
    '2.7':  ('diagrams/fig_2_7_seq_transfer.png',   Inches(6.2)),
    '2.8':  ('diagrams/fig_2_8_seq_verify.png',     Inches(5.5)),
    '2.9':  ('diagrams/fig_2_9_manufacturer.png',   Inches(6.2)),
    '2.10': ('diagrams/fig_2_10_verify.png',        Inches(6.2)),
    '2.11': ('diagrams/fig_2_11_recall.png',        Inches(6.2)),
}

doc = Document(INPUT)

replaced = 0
for para in doc.paragraphs:
    text = para.text.strip()
    # Look for stubs like [МЕСТО ДЛЯ РИСУНКА 2.1]
    if not (text.startswith('[') and 'РИСУНКА' in text.upper()):
        continue

    # Strict regex extraction — never falls back to substring containment.
    match = STUB_RE.search(text)
    if not match:
        print(f'  WARN: cannot parse figure number from stub: {text}')
        continue
    num = match.group(1)
    if num not in FIG_MAP:
        print(f'  SKIP fig {num}: no image mapping (stub will remain as text)')
        continue

    img_path, width = FIG_MAP[num]
    if not os.path.exists(img_path):
        print(f'  SKIP fig {num}: {img_path} not found')
        continue

    # Clear all runs from the paragraph
    p_elem = para._p
    for r_elem in list(p_elem.findall(qn('w:r'))):
        p_elem.remove(r_elem)
    for hl in list(p_elem.findall(qn('w:hyperlink'))):
        p_elem.remove(hl)

    # Set centered alignment on pPr
    pPr = p_elem.find(qn('w:pPr'))
    if pPr is None:
        from docx.oxml import OxmlElement
        pPr = OxmlElement('w:pPr')
        p_elem.insert(0, pPr)
    jc = pPr.find(qn('w:jc'))
    if jc is None:
        from docx.oxml import OxmlElement
        jc = OxmlElement('w:jc')
        pPr.append(jc)
    jc.set(qn('w:val'), 'center')

    # Add image run
    run = para.add_run()
    run.add_picture(img_path, width=width)

    replaced += 1
    print(f'  Embedded fig {num}: {img_path}')

doc.save(OUTPUT)
print(f'\nDone. Replaced {replaced} stubs (of {len(FIG_MAP)} mapped images). Saved {OUTPUT}')
